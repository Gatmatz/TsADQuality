"""
Δημιουργια Word εγγραφου με αποτελεσματα πειραματων.
Output: results/analysis/apotelesmata_peiramatwn.docx
"""
import os
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

current_file_path = os.path.abspath(__file__)
project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(current_file_path))))

RESULTS_BASE = os.path.join(project_root, "results", "experiments")
OUTPUT_PATH = os.path.join(project_root, "results", "analysis", "apotelesmata_peiramatwn.docx")
FONT = 'Times New Roman'
FONT_SIZE = Pt(11)
TABLE_FONT_SIZE = Pt(8)


def add_table(doc, csv_path, columns, labels=None, filters=None):
    if not os.path.exists(csv_path):
        return
    df = pd.read_csv(csv_path)
    if 'model' in df.columns:
        df = df[df['model'] == 'IForest']
    if filters:
        for col, val in filters.items():
            if col in df.columns:
                df = df[df[col] == val]
    cols = [c for c in columns if c in df.columns]
    if not cols or df.empty:
        return
    df_show = df[cols].copy()
    for col in df_show.columns:
        if df_show[col].dtype in ['float64', 'float32']:
            df_show[col] = df_show[col].round(4)
    headers = labels if labels else cols
    table = doc.add_table(rows=1 + len(df_show), cols=len(cols))
    table.style = 'Table Grid'
    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = str(h)
    for i, (_, row) in enumerate(df_show.iterrows()):
        for j, col in enumerate(cols):
            table.rows[i + 1].cells[j].text = str(row[col])
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:
                    r.font.size = TABLE_FONT_SIZE
                    r.font.name = FONT


def add_img(doc, plots_dir, names):
    if not os.path.exists(plots_dir):
        return
    for name in names:
        path = os.path.join(plots_dir, name)
        if os.path.exists(path):
            doc.add_picture(path, width=Inches(4.8))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def p(doc, text):
    par = doc.add_paragraph(text)
    par.style.font.name = FONT
    par.style.font.size = FONT_SIZE


def main():
    os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = FONT_SIZE

    doc.add_heading('Αποτελεσματα Πειραματων', level=0)
    p(doc, 'Dataset: 141 χρονοσειρες (TSB-UAD). Μοντελο: IForest. Baseline AUC-ROC: 0.693.')

    # 1
    doc.add_heading('1. White Noise (SNR)', level=1)
    p(doc, 'Gaussian θορυβος σε ολη τη σειρα. Κατω απο 10dB αρχιζει η πτωση.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "white_noise_snr", "summary.csv"),
        ['snr_db', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['SNR (dB)', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "white_noise_snr", "plots"),
        ['snr_AUC_ROC.png'])

    # 2
    doc.add_heading('2. Missing Values — True Impact', level=1)
    p(doc, 'Burst ελλειψεις. Χαμενα anomalies παιρνουν score=0. block_length = fraction * n / num_bursts.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "missing_true_impact", "summary.csv"),
        ['fraction', 'num_bursts', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Bursts', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "missing_true_impact", "plots"),
        ['true_impact_AUC_ROC.png', 'true_impact_heatmap.png'])

    # 3
    doc.add_heading('3. Missing Values — MCAR vs MNAR (Point)', level=1)
    p(doc, 'Ο Rubin (1976) οριζει τρεις μηχανισμους ελλειπουσων τιμων: '
      'MCAR (τυχαια), MAR (εξαρταται απο αλλες μεταβλητες) και MNAR (εξαρταται απο την ιδια την τιμη). '
      'Εδω συγκρινουμε MCAR με MNAR στο ιδιο ποσοστο missing. '
      'Στο MNAR η πιθανοτητα απωλειας ειναι αναλογη του |z-score|, '
      'δηλαδη οι ακραιες τιμες χανονται πιο συχνα — οπως συμβαινει σε αισθητηρες '
      'που κλεινουν σε υπερφορτωση (sensor saturation). '
      'Αποτελεσμα: στο 20% missing, η MCAR χανει 19% των ανωμαλιων ενω η MNAR χανει 34%.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "missing_mnar", "summary.csv"),
        ['fraction', 'mechanism', 'n_runs', 'mean_pct_anomalies_lost', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Mechanism', 'N', '% Anom. Lost', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "missing_mnar", "plots"),
        ['mnar_AUC_ROC.png', 'mnar_anomalies_lost.png', 'mnar_heatmap.png', 'mnar_corrected_auc.png'])

    # 4
    doc.add_heading('4. Missing Values — MNAR (Burst)', level=1)
    p(doc, 'Ιδια λογικη με το 3 αλλα σε burst αντι point. '
      'Τα bursts τοποθετουνται στις περιοχες με τις πιο ακραιες τιμες. '
      'Η επιδραση ειναι χειροτερη γιατι χανονται ολοκληρες περιοχες ανωμαλιων.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "missing_mnar_burst", "summary.csv"),
        ['fraction', 'num_bursts', 'mechanism', 'n_runs', 'mean_pct_anomalies_lost', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Bursts', 'Mechanism', 'N', '% Anom. Lost', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "missing_mnar_burst", "plots"),
        ['mnar_burst_AUC_ROC.png', 'mnar_burst_heatmap_extreme.png', 'mnar_burst_corrected_auc.png'])

    # 5
    doc.add_heading('5. Gilbert-Elliott', level=1)
    p(doc, 'Μοντελο καναλιου με δυο καταστασεις (Good/Bad) μεσω Markov chain. '
      'Παραγει bursty missing patterns. alpha=G->B, beta=B->G.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "gilbert_elliott_true_impact", "summary.csv"),
        ['alpha', 'beta', 'expected_rate', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['alpha', 'beta', 'Αναμ.%', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "gilbert_elliott_true_impact", "plots"),
        ['ge_true_impact_heatmap.png', 'ge_true_impact_AUC_ROC.png'])

    # 6
    doc.add_heading('6. Freeze (Sensor Stuck)', level=1)
    p(doc, 'Η τιμη κολλαει σε σταθερο επιπεδο (stuck-at-value), '
      'οπως συμβαινει οταν η πηγη δεδομενων σταματα να στελνει ενημερωσεις '
      'και επαναλαμβανεται η τελευταια εγκυρη τιμη (LOCF). '
      'stuck_length = fraction * n / num_stucks.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "freeze", "summary.csv"),
        ['fraction', 'num_stucks', 'mean_stuck_length', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Blocks', 'Μεσο Μηκος', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "freeze", "plots"),
        ['freeze_AUC_ROC.png', 'freeze_heatmap.png', 'freeze_high_vs_low.png'])

    # 7
    doc.add_heading('7. Point Swap', level=1)
    p(doc, 'Ανταλλαγη μεμονωμενων σημειων. Ελαχιστη επιδραση ακομα και στο 40%.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "swap_point", "summary.csv"),
        ['fraction', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "swap_point", "plots"),
        ['swap_point_AUC_ROC.png', 'swap_point_inverted_rate.png', 'swap_point_high_vs_low.png'])

    # 8
    doc.add_heading('8. Segment Swap', level=1)
    p(doc, 'Ανταλλαγη τμηματων. swap_length = fraction * n / (2 * num_swaps).')
    add_table(doc,
        os.path.join(RESULTS_BASE, "swap_segment", "summary.csv"),
        ['fraction', 'num_swaps', 'mean_swap_length', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Swaps', 'Μεσο Μηκος', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "swap_segment", "plots"),
        ['swap_segment_AUC_ROC.png', 'swap_segment_heatmap.png', 'swap_segment_high_vs_low.png'])

    # 9
    doc.add_heading('9. Permutation', level=1)
    p(doc, 'Η σειρα χωριζεται σε N ισα τμηματα και ανακατευεται η σειρα τους. '
      'Ακομα και με N=2 το AUC πεφτει στο 0.499.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "swap_permutation", "summary.csv"),
        ['n_segments', 'mean_segment_length', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['N', 'Μεσο Μηκος', 'N runs', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "swap_permutation", "plots"),
        ['permutation_AUC_ROC_bar.png', 'permutation_corrected_auc.png'])

    # 10
    doc.add_heading('10. Spikes', level=1)
    p(doc, 'Προσθηκη spikes μονο σε κανονικα σημεια. Multiplier σε std. '
      'Στο mult=10x και frac=20% το 86% των αρχειων αντιστρεφονται (AUC<0.5).')
    p(doc, 'Το γραφημα Inverted Detection Rate δειχνει το ποσοστο αρχειων με AUC-ROC < 0.5. '
      'Οταν το AUC πεσει κατω απο 0.5, ο detector εντοπιζει τα spikes αντι τις πραγματικες ανωμαλιες — '
      'δηλαδη η corruption γινεται πιο εμφανης απο το αρχικο σημα. '
      'Στο baseline μονο 21% των αρχειων εχουν AUC<0.5, ενω με 10x spikes στο 20% φτανει στο 86%.')
    add_table(doc,
        os.path.join(RESULTS_BASE, "spikes_normal_only", "summary.csv"),
        ['fraction', 'multiplier', 'n_runs', 'mean_AUC_ROC', 'std_AUC_ROC'],
        ['Fraction', 'Mult', 'N', 'AUC-ROC', 'Std'])
    add_img(doc, os.path.join(RESULTS_BASE, "spikes_normal_only", "plots"),
        ['spikes_normal_AUC_ROC.png', 'spikes_normal_heatmap.png', 'spikes_normal_inverted_rate.png',
         'spikes_normal_corrected_auc.png'])

    # Συνοψη
    doc.add_heading('Συνοψη', level=1)
    ranking = [
        ['Spikes 10x (20%)', '0.329', '-52.6%'],
        ['MNAR extreme burst (20%, nb=5)', '0.347', '-49.9%'],
        ['MNAR extreme point (20%)', '0.432', '-37.7%'],
        ['Permutation (N=2)', '0.499', '-28.0%'],
        ['Spikes 3x (10%)', '0.498', '-28.1%'],
        ['Missing MCAR (20%)', '0.571', '-17.6%'],
        ['SNR = 0dB', '0.620', '-10.6%'],
        ['Missing true impact (10%)', '0.623', '-10.1%'],
        ['Freeze ns=1 (10%)', '0.654', '-5.7%'],
        ['Segment swap ns=1 (10%)', '0.683', '-1.5%'],
        ['Point swap (40%)', '0.678', '-2.3%'],
    ]
    table = doc.add_table(rows=1 + len(ranking), cols=3)
    table.style = 'Table Grid'
    for j, h in enumerate(['Αλλοιωση', 'AUC-ROC', 'Μεταβολη']):
        table.rows[0].cells[j].text = h
    for i, row_data in enumerate(ranking):
        for j, val in enumerate(row_data):
            table.rows[i + 1].cells[j].text = val
    for row in table.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in par.runs:
                    r.font.size = TABLE_FONT_SIZE
                    r.font.name = FONT

    doc.save(OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH}")


if __name__ == '__main__':
    main()
